"""
科研实验数据表:模板导入 → 结构化 schema → 手工录入 → 公式计算 → 导出。

硬约束:
- 模型不参与数值计算;公式由白名单 AST 解释器执行,禁止 eval/任意代码;
- 全程 Decimal,中间不四舍五入,最终输出 ROUND_HALF_UP 到小数点后 4 位;
- Excel 风格公式(=A2*B2、SUM(C2:C10)),为第 2 行书写,逐行自动平移;
- 表格与项目一起按用户存放,复用 research_corpus 的 SQLite。
"""
from __future__ import annotations

import ast
import asyncio
import csv
import io
import json
import os
import re
import secrets
import time
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, getcontext, localcontext
from pathlib import Path
from urllib.parse import quote

from starlette.exceptions import HTTPException
from starlette.requests import Request
from starlette.responses import JSONResponse, Response

import research_corpus as rc
from auth import require_auth
from http_utils import read_json_object
from rag import llm as llm_mod
from rate_limit import SlidingWindowLimiter

getcontext().prec = 28

MAX_TABLES = 100
MAX_ROWS = 200
MAX_COLS = 30
MAX_CELL_CHARS = 200
MAX_FORMULA_CHARS = 500

OUTPUT_QUANT = Decimal("0.0001")
_topic_limiter = SlidingWindowLimiter(12, 3600)

_REF_RE = re.compile(
    r"(?<![A-Za-z])(\$?)([A-Za-z]{1,2})(\$?)(\d+)(?![A-Za-z0-9])")
_RANGE_RE = re.compile(
    r"(?<![A-Za-z])(\$?)([A-Za-z]{1,2})(\$?)(\d+):"
    r"(\$?)([A-Za-z]{1,2})(\$?)(\d+)(?![A-Za-z0-9])")
_COL_KEY_RE = re.compile(r"^[A-Z]{1,2}$")

_ALLOWED_FUNCS = {
    "SUM", "AVERAGE", "AVG", "MIN", "MAX", "ABS", "SQRT", "LOG", "LN",
    "EXP", "POWER", "ROUND", "COUNT",
}
_ALLOWED_BINOPS = (ast.Add, ast.Sub, ast.Mult, ast.Div, ast.Pow, ast.Mod)
_ALLOWED_UNARY = (ast.UAdd, ast.USub)


class FormulaError(Exception):
    """公式语义错误(除零、空值、循环依赖等),按单元格返回。"""


def _col_to_index(key: str) -> int:
    """'A' → 0,'AA' → 26;非法返回 -1。"""
    key = key.upper()
    if not _COL_KEY_RE.match(key):
        return -1
    n = 0
    for ch in key:
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def _index_to_col(index: int) -> str:
    index += 1
    out = ""
    while index > 0:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out


def _dec(value: object) -> Decimal:
    if value is None or value == "":
        raise FormulaError("空值")
    try:
        number = value if isinstance(value, Decimal) else Decimal(str(value).strip())
    except (InvalidOperation, ValueError) as exc:
        raise FormulaError("非数字") from exc
    if not number.is_finite():
        raise FormulaError("非有限数值")
    return number


def _fmt4(value: Decimal) -> str:
    try:
        q = value.quantize(OUTPUT_QUANT, rounding=ROUND_HALF_UP)
    except InvalidOperation as exc:
        raise FormulaError("结果超出范围") from exc
    if q == 0:
        q = abs(q)  # 消除 -0.0000
    return f"{q:.4f}"


def _shift_formula(expr: str, row_index: int) -> str:
    """把为第 2 行书写的公式平移到第 row_index 条数据行。
    数据第 0 行 = Excel 第 2 行;相对引用按偏移平移(带 $ 的行号锚定)。
    Excel 的 ^ 先换成 Python 的 **,再交给 AST。"""
    expr = expr.replace("^", "**")
    offset = row_index

    def range_repl(m: re.Match) -> str:
        c1 = _col_to_index(m.group(2)) + 1
        r1 = int(m.group(4)) + (0 if m.group(3) else offset)  # $ 锚定行不平移
        c2 = _col_to_index(m.group(6)) + 1
        r2 = int(m.group(8)) + (0 if m.group(7) else offset)
        return f"R({c1},{r1},{c2},{r2})"

    def ref_repl(m: re.Match) -> str:
        col = _col_to_index(m.group(2)) + 1
        row = int(m.group(4)) + (0 if m.group(3) else offset)
        return f"C({col},{row})"

    text = _RANGE_RE.sub(range_repl, expr)
    text = _REF_RE.sub(ref_repl, text)
    return text


def _validate_ast(node: ast.AST) -> None:
    if isinstance(node, ast.Expression):
        _validate_ast(node.body)
        return
    if isinstance(node, ast.Constant):
        if isinstance(node.value, bool) or not isinstance(node.value, (int, float)):
            raise FormulaError("公式含非法常量")
        return
    if isinstance(node, ast.BinOp) and isinstance(node.op, _ALLOWED_BINOPS):
        _validate_ast(node.left)
        _validate_ast(node.right)
        return
    if isinstance(node, ast.UnaryOp) and isinstance(node.op, _ALLOWED_UNARY):
        _validate_ast(node.operand)
        return
    if isinstance(node, ast.Call):
        if not isinstance(node.func, ast.Name):
            raise FormulaError("公式含非法调用")
        name = node.func.id.upper()
        if name not in _ALLOWED_FUNCS and name not in ("C", "R"):
            raise FormulaError(f"不支持的函数 {name}")
        if node.keywords:
            raise FormulaError("不支持命名参数")
        for arg in node.args:
            _validate_ast(arg)
        return
    raise FormulaError("公式含不支持的语法")


def _call(name: str, args: list, row: int, rows: list[list[str]],
          col_index: dict[str, int], computed: dict[str, list[Decimal | None]],
          current_col: str) -> object:
    """执行白名单函数;返回 Decimal 或数值列表。"""
    name = name.upper()

    def scalar(item: object) -> Decimal:
        if isinstance(item, list):
            raise FormulaError("该函数不接受区域")
        return _dec(item)

    def numbers(args: list) -> list[Decimal]:
        out: list[Decimal] = []
        for arg in args:
            values = arg if isinstance(arg, list) else [arg]
            for value in values:
                if value is None or value == "":
                    continue  # SUM/AVG 忽略空值,与 Excel 一致
                out.append(_dec(value))
        return out

    if name == "C":
        if len(args) != 2:
            raise FormulaError("单元格引用参数错误")
        c = int(scalar(args[0])) - 1
        # 数据第 0 行 = Excel 第 2 行(第 1 行是表头)
        r = int(scalar(args[1])) - 2
        if c < 0 or r < 0:
            raise FormulaError("引用越界")
        keys = list(col_index.keys())
        if c not in col_index.values():
            raise FormulaError("引用列不存在")
        key = next(k for k, v in col_index.items() if v == c)
        if r < len(rows):
            if key == current_col:
                raise FormulaError("循环引用")
            if key in computed:
                return computed[key][r]
            if c >= len(rows[r]):
                return None
            return rows[r][c]
        raise FormulaError("引用越界")

    if name == "R":
        if len(args) != 4:
            raise FormulaError("区域参数错误")
        c1, r1, c2, r2 = (int(scalar(a)) for a in args)
        out: list[Decimal | None] = []
        for rr in range(min(r1, r2), max(r1, r2) + 1):
            idx = rr - 2
            if not 0 <= idx < len(rows):
                continue
            for cc in range(min(c1, c2), max(c1, c2) + 1):
                col = cc - 1
                if col not in col_index.values():
                    continue
                key = next(k for k, v in col_index.items() if v == col)
                if key == current_col:
                    raise FormulaError("循环引用")
                value = computed[key][idx] if key in computed else (
                    rows[idx][col] if col < len(rows[idx]) else None)
                out.append(value)
        return out

    if name in ("SUM", "COUNT", "AVERAGE", "AVG", "MIN", "MAX"):
        vals = numbers(args)
        if name == "COUNT":
            return Decimal(len(vals))
        if not vals:
            # 没有数据就不产出结果:SUM 也不按 Excel 返回 0,避免把空表算成 0
            raise FormulaError("区域为空")
        if name == "SUM":
            return sum(vals, Decimal(0))
        if name in ("AVERAGE", "AVG"):
            return sum(vals, Decimal(0)) / Decimal(len(vals))
        if name == "MIN":
            return min(vals)
        return max(vals)

    if name == "ABS":
        return abs(scalar(args[0]))
    if name == "SQRT":
        value = scalar(args[0])
        if value < 0:
            raise FormulaError("负数开方")
        return value.sqrt()
    if name == "LOG":
        value = scalar(args[0])
        if value <= 0:
            raise FormulaError("对数定义域错误")
        return value.log10()
    if name == "LN":
        value = scalar(args[0])
        if value <= 0:
            raise FormulaError("对数定义域错误")
        return value.ln()
    if name == "EXP":
        value = scalar(args[0])
        if abs(value) > 1000:
            raise FormulaError("指数过大")
        return value.exp()
    if name == "POWER":
        base = scalar(args[0])
        exponent = scalar(args[1])
        if abs(exponent) > 1000:
            raise FormulaError("指数过大")
        return base ** exponent
    if name == "ROUND":
        value = scalar(args[0])
        digits = int(scalar(args[1]))
        if abs(digits) > 20:
            raise FormulaError("保留位数超出范围")
        return value.quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP)
    raise FormulaError(f"不支持的函数 {name}")


def _eval_node(node: ast.AST, row: int, rows: list[list[str]],
               col_index: dict[str, int], computed: dict[str, list[Decimal | None]],
               current_col: str) -> object:
    if isinstance(node, ast.Expression):
        return _eval_node(node.body, row, rows, col_index, computed, current_col)
    if isinstance(node, ast.Constant):
        return Decimal(str(node.value))
    if isinstance(node, ast.UnaryOp):
        value = _dec(_eval_node(node.operand, row, rows, col_index, computed, current_col))
        return -value if isinstance(node.op, ast.USub) else value
    if isinstance(node, ast.BinOp):
        left = _dec(_eval_node(node.left, row, rows, col_index, computed, current_col))
        right = _dec(_eval_node(node.right, row, rows, col_index, computed, current_col))
        if isinstance(node.op, ast.Add):
            return left + right
        if isinstance(node.op, ast.Sub):
            return left - right
        if isinstance(node.op, ast.Mult):
            return left * right
        if isinstance(node.op, ast.Div):
            if right == 0:
                raise FormulaError("除零")
            return left / right
        if isinstance(node.op, ast.Mod):
            if right == 0:
                raise FormulaError("除零")
            return left % right
        if isinstance(node.op, ast.Pow):
            if abs(right) > 1000:
                raise FormulaError("指数过大")
            try:
                return left ** right
            except (InvalidOperation, ValueError) as exc:
                raise FormulaError("幂运算失败") from exc
    if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
        args = [
            _eval_node(arg, row, rows, col_index, computed, current_col)
            for arg in node.args
        ]
        return _call(node.func.id, args, row, rows, col_index, computed, current_col)
    raise FormulaError("公式求值失败")


def compute(columns: list[dict], rows: list[list[str]], formulas: dict[str, str]) -> dict:
    """返回 {computed: {列: [值...]}, errors: [...], columns: 补全 computed 类型}"""
    col_index: dict[str, int] = {}
    for i, col in enumerate(columns):
        key = str(col.get("key") or _index_to_col(i)).upper()
        col_index[key] = i

    formulas = {str(k).upper(): str(v).strip() for k, v in formulas.items() if str(v).strip()}
    # 依赖排序 + 环检测
    deps: dict[str, set[str]] = {}
    ref_cols_re = re.compile(
        r"(?<![A-Za-z])\$?([A-Za-z]{1,2})\$?\d+(?![A-Za-z0-9])")
    for target, expr in formulas.items():
        deps[target] = {
            m.group(1).upper() for m in ref_cols_re.finditer(expr)
            if m.group(1).upper() in formulas and m.group(1).upper() != target
        }
    order: list[str] = []
    state: dict[str, int] = {}

    def visit(col: str, trail: list[str]) -> None:
        if state.get(col) == 2:
            return
        if state.get(col) == 1:
            cycle = " → ".join(trail + [col])
            raise FormulaError(f"循环依赖:{cycle}")
        state[col] = 1
        for dep in sorted(deps.get(col, set())):
            visit(dep, trail + [col])
        state[col] = 2
        order.append(col)

    try:
        for target in formulas:
            visit(target, [])
    except FormulaError as exc:
        return {"computed": {}, "errors": [{"col": "", "row": -1, "error": str(exc)}],
                "columns": columns}

    computed: dict[str, list[Decimal | None]] = {}
    errors: list[dict] = []
    for target in order:
        expr = formulas[target]
        if not expr.startswith("="):
            expr = "=" + expr
        values: list[Decimal | None] = []
        for row_idx in range(len(rows)):
            try:
                shifted = _shift_formula(expr.lstrip("="), row_idx)
                tree = ast.parse(shifted, mode="eval")
                _validate_ast(tree)
                value = _eval_node(tree, row_idx, rows, col_index, computed, target)
                value = _dec(value)
                _fmt4(value)  # 提前验证可格式化(溢出/NaN)
                values.append(value)
            except FormulaError as exc:
                values.append(None)
                errors.append({"col": target, "row": row_idx, "error": str(exc)})
            except (SyntaxError, ValueError, TypeError) as exc:
                values.append(None)
                errors.append({"col": target, "row": row_idx, "error": f"公式错误:{exc}"})
            except Exception as exc:  # noqa: BLE001 防单格异常拖垮整表计算
                values.append(None)
                errors.append({"col": target, "row": row_idx,
                               "error": f"计算失败:{type(exc).__name__}"})
        computed[target] = values

    display: dict[str, list[str]] = {}
    for target, values in computed.items():
        display[target] = ["" if v is None else _fmt4(v) for v in values]
    inferred = _infer_roles(columns, formulas)
    return {
        "computed": display,
        "errors": errors,
        "columns": [
            {**col, "type": "computed"} if str(col.get("key", "")).upper() in formulas
            else col
            for col in inferred
        ],
    }


# ---------- 统计与图表数据 ----------

def _stat_text(value: Decimal | None) -> str:
    if value is None:
        return ""
    try:
        return _fmt4(value)
    except FormulaError:
        return str(value)


def _percentile(sorted_values: list[Decimal], p: Decimal) -> Decimal | None:
    if not sorted_values:
        return None
    if len(sorted_values) == 1:
        return sorted_values[0]
    rank = Decimal(len(sorted_values) - 1) * p
    lower = int(rank)
    upper = min(lower + 1, len(sorted_values) - 1)
    weight = rank - Decimal(lower)
    return sorted_values[lower] + (sorted_values[upper] - sorted_values[lower]) * weight


def _numeric_columns(
    columns: list[dict], rows: list[list[str]], formulas: dict[str, str],
    computed: dict[str, list[str]], errors: list[dict],
) -> list[dict]:
    """挑出可统计的数值列;文本列只有整列都能解析为数字才算。"""
    out: list[dict] = []
    for i, col in enumerate(columns):
        key = str(col.get("key") or _index_to_col(i)).upper()
        ctype = str(col.get("type") or "text")
        if key in formulas:
            raw = list(computed.get(key) or [])
        else:
            raw = [row[i] if i < len(row) else "" for row in rows]
        values: list[Decimal | None] = []
        parse_failed = False
        for row_index, cell in enumerate(raw):
            text = str(cell).strip()
            if not text:
                values.append(None)
                continue
            try:
                value = Decimal(text)
                # shortcut: 图表数值限制在 1e-323 至 1e307；支持更宽范围时升级图表数值表示。
                if (not value.is_finite() or
                        (value and (value.adjusted() < -323 or value.copy_abs() > Decimal("1e307")))):
                    if ctype in ("number", "computed"):
                        errors.append({"col": key, "row": row_index, "error": "非有限数值或超出支持范围"})
                    values.append(None)
                    parse_failed = True
                    continue
                values.append(value if value else Decimal(0))
            except (InvalidOperation, ValueError):
                values.append(None)
                parse_failed = True
        has_value = any(v is not None for v in values)
        numeric = has_value and (
            ctype in ("number", "computed") or not parse_failed)
        if not numeric:
            continue
        if len(values) < len(rows):
            values.extend([None] * (len(rows) - len(values)))
        out.append({
            "key": key,
            "label": str(col.get("label") or key)[:60],
            "values": values,
        })
    return out


def _pearson(xs: list[Decimal | None], ys: list[Decimal | None]) -> Decimal | None:
    pairs = [(x, y) for x, y in zip(xs, ys) if x is not None and y is not None]
    if len(pairs) < 2:
        return None
    mean_x = sum(x for x, _ in pairs) / Decimal(len(pairs))
    mean_y = sum(y for _, y in pairs) / Decimal(len(pairs))
    dx = [x - mean_x for x, _ in pairs]
    dy = [y - mean_y for _, y in pairs]
    num = sum(a * b for a, b in zip(dx, dy))
    den_x = sum(a * a for a in dx)
    den_y = sum(b * b for b in dy)
    if den_x == 0 or den_y == 0:
        return None
    try:
        return num / (den_x * den_y).sqrt()
    except (InvalidOperation, ValueError):
        return None


def analyze(columns: list[dict], rows: list[list[str]],
            formulas: dict[str, str], row_labels: list[str] | None = None) -> dict:
    """描述性统计 + Pearson 相关 + 绘图序列;全部 Decimal,输出 4 位小数。"""
    result = compute(columns, rows, formulas)
    computed = result["computed"]
    numeric = _numeric_columns(columns, rows, formulas, computed, result["errors"])
    numbers = [v for col in numeric for v in col["values"] if v is not None]
    with localcontext() as ctx:
        if numbers:
            ctx.prec = max(ctx.prec, max(v.adjusted() for v in numbers)
                           - min(v.as_tuple().exponent for v in numbers)
                           + len(str(len(rows))) + 10)
        stats: list[dict] = []
        series: list[dict] = []
        for col in numeric:
            values = [v for v in col["values"] if v is not None]
            ordered = sorted(values)
            n = len(values)
            mean = sum(values) / Decimal(n) if n else None
            stdev = None
            if n >= 2:
                variance = sum((v - mean) ** 2 for v in values) / Decimal(n - 1)
                stdev = variance.sqrt()
            stats.append({
                "key": col["key"], "label": col["label"], "n": n,
                "missing": len(col["values"]) - n,
                "mean": _stat_text(mean),
                "median": _stat_text(_percentile(ordered, Decimal("0.5"))),
                "stdev": _stat_text(stdev),
                "min": _stat_text(ordered[0] if ordered else None),
                "max": _stat_text(ordered[-1] if ordered else None),
                "q1": _stat_text(_percentile(ordered, Decimal("0.25"))),
                "q3": _stat_text(_percentile(ordered, Decimal("0.75"))),
            })
            series.append({
                "key": col["key"], "label": col["label"],
                "values": [None if v is None else float(v) for v in col["values"]],
            })
        correlations: list[dict] = []
        for i in range(len(numeric)):
            for j in range(i + 1, len(numeric)):
                r = _pearson(numeric[i]["values"], numeric[j]["values"])
                if r is None:
                    continue
                correlations.append({
                    "x": numeric[i]["key"], "y": numeric[j]["key"],
                    "r": _stat_text(r),
                })
        labels_in = row_labels if isinstance(row_labels, list) else []
        row_labels_out = [
            str(labels_in[i])[:60] if i < len(labels_in) and labels_in[i] else str(i + 1)
            for i in range(len(rows))
        ]
        return {
            "stats": stats,
            "correlations": correlations,
            "series": series,
            "labels": [str(i + 1) for i in range(len(rows))],
            "rowLabels": row_labels_out,
            "computed": computed,
            "errors": result["errors"],
            "columns": result["columns"],
        }


def context_text(table: dict, max_rows: int = 40) -> str:
    """把数据表压缩成对话上下文:结构、公式、前若干行与计算结果。"""
    columns = table.get("columns") if isinstance(table.get("columns"), list) else []
    rows = table.get("rows") if isinstance(table.get("rows"), list) else []
    formulas = table.get("formulas") if isinstance(table.get("formulas"), dict) else {}
    result = compute(columns, rows, formulas)
    computed = result["computed"]
    labels = table.get("rowLabels") if isinstance(table.get("rowLabels"), list) else []

    def cell(row: list[str], row_i: int, i: int) -> str:
        key = str(columns[i].get("key") or "").upper()
        if key in formulas:
            values = computed.get(key) or []
            text = values[row_i] if row_i < len(values) else ""
        else:
            text = row[i] if i < len(row) else ""
        return str(text)[:40]

    header = ["行"] + [f"{c.get('key')} {str(c.get('label') or '')[:30]}"
                       for c in columns[:12]]
    lines = [
        f"【研究数据表】{str(table.get('title') or '')[:120]}",
        f"规模:{len(rows)} 行 × {len(columns)} 列",
    ]
    source = table.get("source") if isinstance(table.get("source"), dict) else {}
    if source.get("docName"):
        lines.append(f"来源:{str(source.get('docName'))[:120]}"
                     f"{' · 第 ' + str(source.get('page')) + ' 页' if source.get('page') else ''}")
    if formulas:
        lines.append("公式:" + "；".join(
            f"{k}={str(v)[:120]}" for k, v in list(formulas.items())[:12]))
    lines.append("| " + " | ".join(header) + " |")
    lines.append("|" + "---|" * len(header))
    for row_i, row in enumerate(rows[:max_rows]):
        row_label = str(labels[row_i])[:30] if row_i < len(labels) and labels[row_i] else str(row_i + 1)
        lines.append("| " + " | ".join(
            [row_label] + [cell(row, row_i, i) for i in range(min(len(columns), 12))]) + " |")
    if len(rows) > max_rows:
        lines.append(f"（仅列出前 {max_rows} 行,共 {len(rows)} 行）")
    if len(columns) > 12:
        lines.append(f"（表格另有 {len(columns) - 12} 列未列出）")
    stats = analyze(columns, rows, formulas, labels)["stats"]
    if stats:
        lines.append("描述性统计(服务端计算,直接引用,不要自行重算):" + "；".join(
            f"{s['key']} {s['label']} n={s['n']} 均值={s['mean'] or '—'} "
            f"中位数={s['median'] or '—'} 标准差={s['stdev'] or '—'} "
            f"最小={s['min'] or '—'} 最大={s['max'] or '—'}"
            for s in stats[:8]))
    errors = result.get("errors") or []
    if errors:
        lines.append(f"计算错误:{len(errors)} 个单元格,数值结果以能算出的单元格为准。")
    lines.append("以上数据来自用户本项目的数据表,引用数值时必须与表内一致,"
                 "不得改写、编造或用外部资料替代。")
    return "\n".join(lines)[:12000]


def table_context(uid: str, table_id: str) -> str | None:
    table = _get_table(uid, table_id)
    return context_text(table) if table else None


# ---------- 导入 / 存储 ----------

def _clean_cell(value: object) -> str:
    text = "" if value is None else str(value).strip()
    return text[:MAX_CELL_CHARS]


def _parse_csv(data: bytes) -> list[list[str]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("gb18030", errors="replace")
    return [[_clean_cell(c) for c in row] for row in csv.reader(io.StringIO(text))]


def _parse_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        out: list[list[str]] = []
        for row in ws.iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLS + 1, values_only=True):
            out.append([_clean_cell(c) for c in row[:MAX_COLS]])
        return out
    finally:
        wb.close()


def _parse_docx(data: bytes) -> list[list[str]]:
    """DOCX 里的第一个表格;没有表格直接报错(不做段落猜测)。"""
    from docx import Document

    doc = Document(io.BytesIO(data))
    if not doc.tables:
        raise ValueError("DOCX 中没有表格")
    table = doc.tables[0]
    return [
        [_clean_cell(cell.text) for cell in row.cells[:MAX_COLS]]
        for row in table.rows[:MAX_ROWS + 1]
    ]


def _table_from_grid(grid: list[list[str]], title: str) -> dict:
    # XLSX 只读模式会把空列一起读出来:先裁掉每行尾部空单元格
    trimmed: list[list[str]] = []
    for row in grid:
        row = list(row)
        while row and row[-1] == "":
            row.pop()
        trimmed.append(row)
    grid = [row for row in trimmed if any(cell for cell in row)]
    if not grid:
        raise ValueError("表格为空")
    if len(grid) > MAX_ROWS + 1:
        raise ValueError(f"行数超过上限 {MAX_ROWS}")
    width = min(max(len(row) for row in grid), MAX_COLS)
    if width == 0:
        raise ValueError("表格为空")
    headers = grid[0][:width]
    body = grid[1:MAX_ROWS + 1]
    columns = []
    for i in range(width):
        label = headers[i] or _index_to_col(i)
        values = [row[i] if i < len(row) else "" for row in body]
        numeric = bool(values) and all(
            not v or _is_number(v) for v in values)
        columns.append({"key": _index_to_col(i), "label": label[:60],
                        "type": "number" if numeric else "text"})
    rows = [[row[i] if i < len(row) else "" for i in range(width)] for row in body]
    return {
        "title": title[:120] or "实验数据表",
        "columns": columns,
        "rows": rows,
        "formulas": {},
    }


def _is_number(text: str) -> bool:
    try:
        Decimal(str(text).strip())
        return True
    except (InvalidOperation, ValueError):
        return False


def _connect(uid: str):
    conn = rc._connect(uid)
    conn.execute(
        "CREATE TABLE IF NOT EXISTS tables("
        "id TEXT PRIMARY KEY, project_id TEXT, title TEXT NOT NULL,"
        "columns_json TEXT NOT NULL, rows_json TEXT NOT NULL,"
        "formulas_json TEXT NOT NULL, created_at REAL NOT NULL, updated_at REAL NOT NULL,"
        "row_labels_json TEXT NOT NULL DEFAULT '[]',"
        "merges_json TEXT NOT NULL DEFAULT '[]',"
        "col_widths_json TEXT NOT NULL DEFAULT '[]',"
        "row_heights_json TEXT NOT NULL DEFAULT '[]',"
        "source_json TEXT NOT NULL DEFAULT '{}',"
        "row_sources_json TEXT NOT NULL DEFAULT '[]')"
    )
    columns = {row[1] for row in conn.execute("PRAGMA table_info(tables)")}
    for column, decl in (
        ("row_labels_json", "TEXT NOT NULL DEFAULT '[]'"),
        ("merges_json", "TEXT NOT NULL DEFAULT '[]'"),
        ("col_widths_json", "TEXT NOT NULL DEFAULT '[]'"),
        ("row_heights_json", "TEXT NOT NULL DEFAULT '[]'"),
        ("source_json", "TEXT NOT NULL DEFAULT '{}'"),
        ("row_sources_json", "TEXT NOT NULL DEFAULT '[]'"),
    ):
        if column not in columns:
            conn.execute(f"ALTER TABLE tables ADD COLUMN {column} {decl}")
    return conn


def _json_list(raw: object) -> list:
    try:
        value = json.loads(raw) if isinstance(raw, str) else raw
        return value if isinstance(value, list) else []
    except (TypeError, json.JSONDecodeError):
        return []


def _json_dict(raw: object) -> dict:
    try:
        value = json.loads(raw) if isinstance(raw, str) else raw
        return value if isinstance(value, dict) else {}
    except (TypeError, json.JSONDecodeError):
        return {}


def _clean_source(raw: object) -> dict:
    """数据表来源:只保留可追溯字段,不接受任意对象膨胀存储。"""
    if not isinstance(raw, dict):
        return {}
    doc_id = str(raw.get("docId") or "").strip()[:64]
    if not doc_id:
        return {}
    try:
        page = int(raw.get("page") or 0)
    except (TypeError, ValueError):
        page = 0
    source = {
        "docId": doc_id,
        "docName": str(raw.get("docName") or "").strip()[:200],
        "page": max(0, min(page, 100000)),
        "quote": str(raw.get("quote") or "").strip()[:300],
        "kind": "paper",
    }
    return source


def _clean_row_sources(raw: object, n_rows: int) -> list[dict]:
    """每行一个来源(多文献对比表用);空行保留 {} 以维持行对齐。"""
    sources: list[dict] = [{} for _ in range(n_rows)]
    if not isinstance(raw, list):
        return sources
    for i, item in enumerate(raw[:n_rows]):
        sources[i] = _clean_source(item)
    return sources


_IDENT_RE = re.compile(
    r"(编号|序号|样本|样品|组别|分组|日期|时间点|(?<![A-Za-z])(?:id|sample|group|index|no\.?|date)(?![A-Za-z]))",
    re.IGNORECASE)


def _infer_roles(columns: list[dict], formulas: dict[str, str]) -> list[dict]:
    """按公式依赖补列角色:目标=计算,被引用=实测;已有合法角色不覆盖。"""
    targets = {str(k).upper() for k in formulas}
    refs: set[str] = set()
    for expr in formulas.values():
        for match in re.finditer(
                r"(?<![A-Za-z])\$?([A-Za-z]{1,2})\$?\d+(?![A-Za-z0-9])",
                str(expr)):
            refs.add(match.group(1).upper())
        for match in _RANGE_RE.finditer(str(expr)):
            start = _col_to_index(match.group(2))
            end = _col_to_index(match.group(6))
            refs.update(_index_to_col(i) for i in range(min(start, end), max(start, end) + 1))
    out = []
    for col in columns:
        item = dict(col)
        key = str(item.get("key") or "").upper()
        role = str(item.get("role") or "")
        if key in targets:
            item["role"] = "calculated"
        elif role not in ("measured", "calculated", "identifier", "unit", "note"):
            if key in refs:
                item["role"] = "measured"
            elif _IDENT_RE.search(str(item.get("label") or "")):
                item["role"] = "identifier"
        out.append(item)
    return out


def _row_to_table(row: tuple) -> dict:
    columns = json.loads(row[3])
    formulas = json.loads(row[5])
    return {
        "id": row[0], "projectId": row[1], "title": row[2],
        "columns": _infer_roles(columns, formulas), "rows": json.loads(row[4]),
        "formulas": formulas, "createdAt": row[6], "updatedAt": row[7],
        "rowLabels": _json_list(row[8] if len(row) > 8 else []),
        "merges": _json_list(row[9] if len(row) > 9 else []),
        "colWidths": _json_list(row[10] if len(row) > 10 else []),
        "rowHeights": _json_list(row[11] if len(row) > 11 else []),
        "source": _json_dict(row[12] if len(row) > 12 else None),
        "rowSources": _json_list(row[13] if len(row) > 13 else []),
    }


_TABLE_SELECT = (
    "SELECT id, project_id, title, columns_json, rows_json, formulas_json,"
    " created_at, updated_at, row_labels_json, merges_json, col_widths_json,"
    " row_heights_json, source_json, row_sources_json FROM tables"
)


def _list_tables(uid: str, project_id: str | None) -> list[dict]:
    conn = _connect(uid)
    try:
        sql = _TABLE_SELECT
        params: tuple = ()
        if project_id:
            sql += " WHERE project_id = ?"
            params = (project_id,)
        sql += " ORDER BY updated_at DESC"
        out = []
        for row in conn.execute(sql, params):
            table = _row_to_table(row)
            out.append({
                "id": table["id"], "projectId": table["projectId"],
                "title": table["title"], "columns": len(table["columns"]),
                "rows": len(table["rows"]), "updatedAt": table["updatedAt"],
            })
        return out
    finally:
        conn.close()


def _get_table(uid: str, table_id: str) -> dict | None:
    conn = _connect(uid)
    try:
        row = conn.execute(
            _TABLE_SELECT + " WHERE id = ?", (table_id,)).fetchone()
        return _row_to_table(row) if row else None
    finally:
        conn.close()


def _clean_merges(raw: object, n_rows: int, n_cols: int) -> list[dict]:
    if not isinstance(raw, list):
        return []
    merges: list[dict] = []
    occupied: set[tuple[int, int]] = set()
    for item in raw[:500]:
        if not isinstance(item, dict):
            continue
        try:
            r1 = int(item.get("r1")); c1 = int(item.get("c1"))
            r2 = int(item.get("r2")); c2 = int(item.get("c2"))
        except (TypeError, ValueError):
            continue
        r1, r2 = sorted((r1, r2))
        c1, c2 = sorted((c1, c2))
        if not (0 <= r1 <= r2 < n_rows and 0 <= c1 <= c2 < n_cols):
            raise ValueError("合并区域超出表格范围")
        if r1 == r2 and c1 == c2:
            continue
        cells = {(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)}
        if occupied & cells:
            raise ValueError("合并区域不能重叠")
        occupied |= cells
        merges.append({"r1": r1, "c1": c1, "r2": r2, "c2": c2})
    return merges


def _clean_sizes(raw: object, n: int, default: int, low: int, high: int) -> list[int]:
    values: list[int] = []
    if isinstance(raw, list):
        for item in raw[:n]:
            try:
                values.append(max(low, min(int(item), high)))
            except (TypeError, ValueError):
                values.append(default)
    while len(values) < n:
        values.append(default)
    return values


def _validate_table_payload(
    body: dict,
) -> tuple[list[dict], list[list[str]], dict[str, str],
           list[str], list[dict], list[int], list[int], list[dict]]:
    columns = body.get("columns")
    rows = body.get("rows")
    formulas = body.get("formulas", {})
    if not isinstance(columns, list) or not isinstance(rows, list):
        raise ValueError("columns/rows 必须是数组")
    if not 1 <= len(columns) <= MAX_COLS:
        raise ValueError(f"列数必须在 1~{MAX_COLS}")
    if len(rows) > MAX_ROWS:
        raise ValueError(f"行数不能超过 {MAX_ROWS}")
    clean_cols: list[dict] = []
    for i, col in enumerate(columns):
        if not isinstance(col, dict):
            raise ValueError("列定义格式错误")
        key = str(col.get("key") or _index_to_col(i)).upper()
        if not _COL_KEY_RE.match(key):
            key = _index_to_col(i)
        label = str(col.get("label") or key)[:60]
        ctype = str(col.get("type") or "text")
        if ctype not in ("number", "text", "computed"):
            ctype = "text"
        item = {"key": key, "label": label, "type": ctype}
        role = str(col.get("role") or "")
        if role in ("measured", "calculated", "identifier", "unit", "note"):
            item["role"] = role
        unit = str(col.get("unit") or "").strip()[:30]
        if unit:
            item["unit"] = unit
        clean_cols.append(item)
    clean_rows: list[list[str]] = []
    for row in rows:
        if not isinstance(row, list):
            raise ValueError("行格式错误")
        clean_rows.append([_clean_cell(v) for v in row[:len(clean_cols)]])
    if not isinstance(formulas, dict):
        raise ValueError("formulas 必须是对象")
    clean_formulas: dict[str, str] = {}
    for key, expr in formulas.items():
        target = str(key).upper()
        if not _COL_KEY_RE.match(target):
            continue
        if not any(c["key"] == target for c in clean_cols):
            continue
        text = str(expr).strip()[:MAX_FORMULA_CHARS]
        if text:
            clean_formulas[target] = text
    raw_labels = body.get("rowLabels")
    row_labels = [str(v)[:60] for v in raw_labels[:len(clean_rows)]] \
        if isinstance(raw_labels, list) else []
    while len(row_labels) < len(clean_rows):
        row_labels.append("")
    merges = _clean_merges(body.get("merges"), len(clean_rows), len(clean_cols))
    col_widths = _clean_sizes(body.get("colWidths"), len(clean_cols), 120, 40, 600)
    row_heights = _clean_sizes(body.get("rowHeights"), len(clean_rows), 30, 24, 300)
    row_sources = _clean_row_sources(body.get("rowSources"), len(clean_rows))
    return (clean_cols, clean_rows, clean_formulas,
            row_labels, merges, col_widths, row_heights, row_sources)


def _save_table(uid: str, table_id: str, body: dict) -> dict:
    (columns, rows, formulas, row_labels, merges,
     col_widths, row_heights, row_sources) = _validate_table_payload(body)
    title = str(body.get("title") or "实验数据表")[:120]
    now = time.time()
    conn = _connect(uid)
    try:
        existing = conn.execute(
            "SELECT source_json, row_sources_json FROM tables WHERE id = ?",
            (table_id,)).fetchone()
        if existing is None:
            raise KeyError(table_id)
        # 旧客户端不带来源字段时保留原值,避免自动保存把溯源信息抹掉
        source = (_clean_source(body.get("source"))
                  if "source" in body else _json_dict(existing[0]))
        if "rowSources" not in body:
            row_sources = _clean_row_sources(
                _json_list(existing[1]), len(rows))
        cur = conn.execute(
            "UPDATE tables SET title=?, columns_json=?, rows_json=?, formulas_json=?,"
            " row_labels_json=?, merges_json=?, col_widths_json=?, row_heights_json=?,"
            " source_json=?, row_sources_json=?, updated_at=? WHERE id=?",
            (title, json.dumps(columns, ensure_ascii=False),
             json.dumps(rows, ensure_ascii=False),
             json.dumps(formulas, ensure_ascii=False),
             json.dumps(row_labels, ensure_ascii=False),
             json.dumps(merges, ensure_ascii=False),
             json.dumps(col_widths), json.dumps(row_heights),
             json.dumps(source, ensure_ascii=False),
             json.dumps(row_sources, ensure_ascii=False), now, table_id))
        if cur.rowcount == 0:
            raise KeyError(table_id)
        conn.commit()
    finally:
        conn.close()
    return {"id": table_id, "title": title, "columns": columns, "rows": rows,
            "formulas": formulas, "rowLabels": row_labels, "merges": merges,
            "colWidths": col_widths, "rowHeights": row_heights,
            "source": source, "rowSources": row_sources, "updatedAt": now}


def _create_table(uid: str, body: dict, project_id: str | None) -> dict:
    (columns, rows, formulas, row_labels, merges,
     col_widths, row_heights, row_sources) = _validate_table_payload(body)
    title = str(body.get("title") or "实验数据表")[:120]
    source = _clean_source(body.get("source"))
    table_id = f"t-{secrets.token_hex(8)}"
    now = time.time()
    conn = _connect(uid)
    try:
        if project_id:
            found = conn.execute("SELECT id FROM projects WHERE id=?", (project_id,)).fetchone()
            if not found:
                raise ValueError("研究项目不存在")
        count = conn.execute("SELECT COUNT(*) FROM tables").fetchone()[0]
        if count >= MAX_TABLES:
            raise ValueError(f"数据表数量达到上限 {MAX_TABLES}")
        conn.execute(
            "INSERT INTO tables(id, project_id, title, columns_json, rows_json,"
            " formulas_json, created_at, updated_at, row_labels_json, merges_json,"
            " col_widths_json, row_heights_json, source_json, row_sources_json)"
            " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (table_id, project_id, title, json.dumps(columns, ensure_ascii=False),
             json.dumps(rows, ensure_ascii=False),
             json.dumps(formulas, ensure_ascii=False), now, now,
             json.dumps(row_labels, ensure_ascii=False),
             json.dumps(merges, ensure_ascii=False),
             json.dumps(col_widths), json.dumps(row_heights),
             json.dumps(source, ensure_ascii=False),
             json.dumps(row_sources, ensure_ascii=False)))
        conn.commit()
    finally:
        conn.close()
    return {"id": table_id, "projectId": project_id, "title": title,
            "columns": columns, "rows": rows, "formulas": formulas,
            "rowLabels": row_labels, "merges": merges,
            "colWidths": col_widths, "rowHeights": row_heights,
            "source": source, "rowSources": row_sources,
            "createdAt": now, "updatedAt": now}


def tables_for_export(uid: str) -> list[dict]:
    """账号导出用:全部数据表(含草稿与公式)。"""
    if not rc._db_path(uid).exists():
        return []
    conn = _connect(uid)
    try:
        return [
            _row_to_table(row)
            for row in conn.execute(_TABLE_SELECT + " ORDER BY updated_at DESC")
        ]
    finally:
        conn.close()


def _delete_table(uid: str, table_id: str, *, hard: bool = False) -> bool:
    conn = _connect(uid)
    try:
        conn.execute("DELETE FROM formulas WHERE table_id = ?", (table_id,))
        cur = conn.execute("DELETE FROM tables WHERE id = ?", (table_id,))
        conn.commit()
        if cur.rowcount <= 0:
            return False
        if hard:
            # 彻底清除:回收空闲页并截断 WAL;之后新写入不会复用旧页
            try:
                conn.execute("PRAGMA wal_checkpoint(TRUNCATE)")
                conn.execute("VACUUM")
            except Exception as exc:  # noqa: BLE001 清理失败不影响删除结果
                print(f"[tables] 彻底清除后压缩失败(忽略): {exc}")
        return True
    finally:
        conn.close()


def _import_table(uid: str, name: str, data: bytes, project_id: str | None) -> dict:
    lower = name.lower()
    if lower.endswith(".xlsx"):
        grid = _parse_xlsx(data)
    elif lower.endswith(".docx"):
        grid = _parse_docx(data)
    elif lower.endswith((".csv", ".txt")):
        grid = _parse_csv(data)
    else:
        raise ValueError("只支持 XLSX、DOCX 与 CSV 模板")
    title = Path(name).stem or "实验数据表"
    body = _table_from_grid(grid, title)
    return _create_table(uid, body, project_id)


# ---------- 导出 ----------

def _table_rows_with_computed(table: dict) -> tuple[list[str], list[list[str]]]:
    result = compute(table["columns"], table["rows"], table["formulas"])
    headers = ["名称"] + [str(c.get("label") or c.get("key")) for c in table["columns"]]
    labels = table.get("rowLabels") if isinstance(table.get("rowLabels"), list) else []
    rows: list[list[str]] = []
    for i, row in enumerate(table["rows"]):
        out = [str(labels[i]) if i < len(labels) else ""] + list(row)
        for target, values in result["computed"].items():
            idx = next((j for j, c in enumerate(table["columns"])
                        if str(c.get("key")).upper() == target), None)
            if idx is None:
                continue
            while len(out) <= idx + 1:
                out.append("")
            out[idx + 1] = values[i] if i < len(values) else ""
        rows.append(out)
    return headers, rows


def _csv_safe(value: str) -> str:
    """CSV 公式注入防护:文本以 = + - @ 开头时加单引号。
    纯数字(含负号)保持原样,否则导出后无法参与计算。"""
    if value and value[0] in "=+-@":
        if _is_number(value):
            return value
        return "'" + value
    return value


def _export_csv(table: dict) -> bytes:
    headers, rows = _table_rows_with_computed(table)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([_csv_safe(h) for h in headers])
    for row in rows:
        writer.writerow([_csv_safe(c) for c in row])
    return buf.getvalue().encode("utf-8-sig")


def _export_xlsx(table: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    headers, rows = _table_rows_with_computed(table)
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(headers)
    for row in rows:
        ws.append([float(v) if _is_number(v) else _csv_safe(v) for v in row])
    for idx, col in enumerate(table["columns"], 2):
        if str(col.get("key")).upper() in table["formulas"]:
            for cell in ws.iter_cols(min_col=idx, max_col=idx, min_row=2):
                for c in cell:
                    c.number_format = "0.0000"
    for merge in table.get("merges") or []:
        if not isinstance(merge, dict):
            continue
        try:
            ws.merge_cells(
                start_row=int(merge["r1"]) + 2, start_column=int(merge["c1"]) + 2,
                end_row=int(merge["r2"]) + 2, end_column=int(merge["c2"]) + 2)
        except (KeyError, TypeError, ValueError):
            continue
    widths = table.get("colWidths") or []
    for i, width in enumerate(widths, 1):
        try:
            ws.column_dimensions[get_column_letter(i)].width = max(6, int(width) / 7)
        except (TypeError, ValueError):
            continue
    heights = table.get("rowHeights") or []
    for i, height in enumerate(heights, 2):
        try:
            ws.row_dimensions[i].height = max(12, int(height) * 0.75)
        except (TypeError, ValueError):
            continue
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- 本地打印 ----------

def _print_dir(uid: str) -> Path:
    path = Path(rc.DATA_DIR) / "print" / rc._safe_uid(uid)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _safe_filename(title: str) -> str:
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", title or "").strip(" .")
    return (name or "table")[:40]


def _prune_print_files(folder: Path, keep: int = 20) -> None:
    files = sorted(
        folder.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    for path in files[keep:]:
        try:
            path.unlink()
        except OSError:
            continue


def _write_print_copy(uid: str, table: dict) -> Path:
    folder = _print_dir(uid)
    name = f"{_safe_filename(table.get('title') or '')}-{time.strftime('%Y%m%d-%H%M%S')}.xlsx"
    path = folder / name
    path.write_bytes(_export_xlsx(table))
    _prune_print_files(folder)
    return path


def _open_in_office(path: Path) -> bool:
    """用系统默认表格程序打开(WPS/Excel),打印选项沿用该程序。"""
    startfile = getattr(os, "startfile", None)
    if startfile is None:
        return False
    try:
        startfile(str(path))
        return True
    except OSError:
        return False


def _print_allowed(request: Request) -> bool:
    """只允许本机页面触发打开本地程序;部署到远端时返回 403。"""
    if os.getenv("MG_PRINT_OPEN") == "1":
        return True
    host = (request.headers.get("host") or "").strip().lower()
    if host.startswith("["):
        name = host[1:host.find("]")] if "]" in host else host
    else:
        name = host.split(":")[0]
    return name in {"localhost", "127.0.0.1", "::1"} or name.endswith(".localhost")


# ---------- HTTP ----------

async def _read_limited(request: Request) -> bytes:
    limit = 5 * 1024 * 1024
    chunks: list[bytes] = []
    total = 0
    async for chunk in request.stream():
        total += len(chunk)
        if total > limit:
            raise HTTPException(413, "模板文件超过 5MB 上限")
        chunks.append(chunk)
    return b"".join(chunks)


async def list_tables(request: Request) -> JSONResponse:
    rec = require_auth(request)
    uid = str(rec.get("uid") or "")
    project_id = request.query_params.get("projectId") or None
    tables = await asyncio.to_thread(_list_tables, uid, project_id)
    return JSONResponse({"tables": tables}, headers={"Cache-Control": "no-store"})


async def import_table(request: Request) -> JSONResponse:
    rec = require_auth(request)
    uid = str(rec.get("uid") or "")
    name = request.query_params.get("name", "").strip()
    project_id = request.query_params.get("projectId") or None
    data = await _read_limited(request)
    if not data:
        raise HTTPException(400, "empty file")
    try:
        table = await asyncio.to_thread(_import_table, uid, name, data, project_id)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    return JSONResponse({"table": table})


async def get_table(request: Request) -> JSONResponse:
    rec = require_auth(request)
    uid = str(rec.get("uid") or "")
    table_id = str(request.path_params.get("table_id") or "")
    table = await asyncio.to_thread(_get_table, uid, table_id)
    if table is None:
        raise HTTPException(404, "not found")
    return JSONResponse({"table": table}, headers={"Cache-Control": "no-store"})


async def save_table(request: Request) -> JSONResponse:
    rec = require_auth(request)
    uid = str(rec.get("uid") or "")
    table_id = str(request.path_params.get("table_id") or "")
    body = await read_json_object(request)
    try:
        table = await asyncio.to_thread(_save_table, uid, table_id, body)
    except KeyError as exc:
        raise HTTPException(404, "not found") from exc
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    return JSONResponse({"table": table})


async def create_table(request: Request) -> JSONResponse:
    rec = require_auth(request)
    uid = str(rec.get("uid") or "")
    body = await read_json_object(request)
    project_id = body.get("projectId")
    project_id = str(project_id) if isinstance(project_id, str) and project_id else None
    try:
        table = await asyncio.to_thread(_create_table, uid, body, project_id)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    return JSONResponse({"table": table})


async def delete_table(request: Request) -> JSONResponse:
    rec = require_auth(request)
    uid = str(rec.get("uid") or "")
    table_id = str(request.path_params.get("table_id") or "")
    hard = request.query_params.get("hard", "").lower() in {"1", "true", "yes"}
    ok = await asyncio.to_thread(_delete_table, uid, table_id, hard=hard)
    if not ok:
        raise HTTPException(404, "not found")
    return JSONResponse({"ok": True, "hard": hard})


async def compute_table(request: Request) -> JSONResponse:
    """「完成」:保存当前草稿并计算;返回计算结果与逐格错误。"""
    rec = require_auth(request)
    uid = str(rec.get("uid") or "")
    table_id = str(request.path_params.get("table_id") or "")
    body = await read_json_object(request)
    try:
        saved = await asyncio.to_thread(_save_table, uid, table_id, body)
    except KeyError as exc:
        raise HTTPException(404, "not found") from exc
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    result = await asyncio.to_thread(
        compute, saved["columns"], saved["rows"], saved["formulas"])
    return JSONResponse({"table": saved, **result})


async def validate_formula(request: Request) -> JSONResponse:
    """反算校验:用候选公式重算目标列,与已有数值比对;不落库。"""
    rec = require_auth(request)
    uid = str(rec.get("uid") or "")
    table_id = str(request.path_params.get("table_id") or "")
    table = await asyncio.to_thread(_get_table, uid, table_id)
    if table is None:
        raise HTTPException(404, "not found")
    body = await read_json_object(request)
    column_key = str(body.get("columnKey") or "").strip().upper()
    formula = str(body.get("formula") or "").strip()[:MAX_FORMULA_CHARS]
    if not column_key or not formula:
        raise HTTPException(400, "缺少 columnKey 或 formula")
    columns = table.get("columns") or []
    target_index = next(
        (i for i, col in enumerate(columns)
         if str(col.get("key") or "").upper() == column_key), None)
    if target_index is None:
        raise HTTPException(400, "目标列不存在")
    rows = table.get("rows") or []
    result = await asyncio.to_thread(
        compute, columns, rows, {column_key: formula})
    predicted = result["computed"].get(column_key, [])
    matched = 0
    compared = 0
    max_error = Decimal(0)
    samples: list[dict] = []
    for index, row in enumerate(rows):
        actual_text = str(row[target_index] if target_index < len(row) else "").strip()
        if not actual_text or index >= len(predicted) or not predicted[index]:
            continue
        try:
            actual = _dec(actual_text)
            value = _dec(predicted[index])
            if actual and (actual.adjusted() < -323 or actual.copy_abs() > Decimal("1e307")):
                raise FormulaError("数值超出支持范围")
        except FormulaError as exc:
            result["errors"].append({"col": column_key, "row": index, "error": str(exc)})
            continue
        compared += 1
        error = abs(actual - value)
        if error > max_error:
            max_error = error
        tolerance = max(abs(actual) * Decimal("0.000001"), Decimal("0.00005"))
        ok = error <= tolerance
        if ok:
            matched += 1
        if len(samples) < 5:
            samples.append({
                "row": index,
                "actual": _stat_text(actual),
                "predicted": predicted[index],
                "ok": ok,
            })
    return JSONResponse({
        "columnKey": column_key,
        "matched": matched,
        "compared": compared,
        "maxError": _stat_text(max_error) if compared else "",
        "samples": samples,
        "errors": result["errors"][:5],
    }, headers={"Cache-Control": "no-store"})


_FUNC_MAP = {
    "sqrt": "SQRT", "abs": "ABS", "ln": "LN", "log": "LOG", "exp": "EXP",
    "round": "ROUND", "sum": "SUM", "avg": "AVERAGE", "mean": "AVERAGE",
    "min": "MIN", "max": "MAX", "power": "POWER",
}


def _compile_table_formula(expr: str, mapping: dict[str, str]) -> str:
    """符号公式 → 表格 Excel 公式;变量必须全部映射。"""
    text = str(expr or "").strip()
    if "=" in text:
        text = text[text.rfind("=") + 1:]
    text = (text.replace("×", "*").replace("÷", "/")
            .replace("−", "-").replace("–", "-").replace("·", "*")
            .replace("²", "^2").replace("³", "^3"))
    for name, target in _FUNC_MAP.items():
        text = re.sub(rf"\b{name}\s*\(", f"{target}(", text, flags=re.IGNORECASE)
    if any(not variable or not key for variable, key in mapping.items()):
        return ""
    if mapping:
        patterns = []
        for variable in sorted(mapping, key=len, reverse=True):
            escaped = re.escape(variable)
            if re.fullmatch(r"[A-Za-z0-9_]+", variable):
                escaped = rf"(?<![A-Za-z0-9_]){escaped}(?![A-Za-z0-9_])"
            patterns.append(escaped)
        functions = "|".join(map(re.escape, set(_FUNC_MAP.values())))
        tokens = re.compile(rf"(\b(?:{functions})\s*\()|({'|'.join(patterns)})")
        text = tokens.sub(lambda match: match.group(1) or f"{mapping[match.group(2)]}2", text)
    return "=" + text


def _match_formula_candidates(table: dict, saved: list[dict]) -> list[dict]:
    """按表头、变量映射和反算结果给已保存公式打分,返回 Top-3。"""
    columns = table.get("columns") or []
    rows = table.get("rows") or []
    candidates = []
    for item in saved[:50]:
        target = str(item.get("target") or "").strip()
        expr = str(item.get("expr") or "").strip()
        variables = item.get("variables") if isinstance(item.get("variables"), list) else []
        suggested = item.get("variableMap") if isinstance(item.get("variableMap"), dict) else {}
        units = item.get("units") if isinstance(item.get("units"), dict) else {}
        if not target or not expr:
            continue
        target_col = next(
            (c for c in columns if str(c.get("label") or "") == target), None)
        target_key = str(target_col.get("key") or "") if target_col else ""
        target_index = (next((i for i, c in enumerate(columns)
                              if c is target_col), -1)
                        if target_col else -1)
        mapping: dict[str, str] = {}
        matched_vars = 0
        score = 0.0
        unit_conflicts: list[dict] = []
        for variable in variables:
            hint = str(suggested.get(variable) or "").strip()
            hit = next(
                (c for c in columns
                 if hint and (str(c.get("label") or "") == hint
                              or str(c.get("key") or "") == hint)), None)
            if hit is None:
                hit = next(
                    (c for c in columns
                     if variable and (variable in str(c.get("label") or "")
                                      or str(c.get("label") or "") in variable)), None)
            if hit:
                mapping[variable] = str(hit.get("key") or "")
                matched_vars += 1
                expected = str(units.get(variable) or "").strip()
                actual_unit = str(hit.get("unit") or "").strip()
                if expected and actual_unit:
                    if expected.replace(" ", "") == actual_unit.replace(" ", ""):
                        score += 0.5
                    else:
                        score -= 1
                        unit_conflicts.append({
                            "variable": variable, "expected": expected, "actual": actual_unit,
                        })
        score += matched_vars * 2
        if target_col:
            score += 2
            expected = str(units.get(target) or "").strip()
            actual_unit = str(target_col.get("unit") or "").strip()
            if expected and actual_unit:
                if expected.replace(" ", "") == actual_unit.replace(" ", ""):
                    score += 0.5
                else:
                    score -= 1
                    unit_conflicts.append({
                        "variable": target, "expected": expected, "actual": actual_unit,
                    })
            role = str(target_col.get("role") or "")
            if role == "calculated":
                score += 2
            if target_key:
                empty = all(
                    target_index >= len(row)
                    or not str(row[target_index]).strip() for row in rows)
                if empty:
                    score += 1
        if variables and matched_vars < len(variables):
            score -= 2
        formula = (_compile_table_formula(expr, mapping)
                   if mapping and matched_vars == len(variables) else "")
        validation = None
        if formula and target_col:
            try:
                result = compute(columns, rows, {
                    target_key.upper(): formula})
                predicted = result["computed"].get(target_key.upper(), [])
                matched = compared = unverified = 0
                max_error = Decimal(0)
                samples: list[dict] = []
                for index, row in enumerate(rows):
                    actual_text = str(
                        row[target_index] if 0 <= target_index < len(row) else "").strip()
                    if not actual_text:
                        continue
                    if index >= len(predicted) or not predicted[index]:
                        unverified += 1
                        continue
                    try:
                        actual = _dec(actual_text)
                        value = _dec(predicted[index])
                        if actual and (actual.adjusted() < -323 or actual.copy_abs() > Decimal("1e307")):
                            unverified += 1
                            continue
                    except FormulaError:
                        unverified += 1
                        continue
                    compared += 1
                    error = abs(actual - value)
                    max_error = max(max_error, error)
                    ok = error <= max(abs(actual) * Decimal("0.000001"), Decimal("0.00005"))
                    if ok:
                        matched += 1
                    if len(samples) < 3:
                        samples.append({
                            "row": index,
                            "actual": _stat_text(actual),
                            "predicted": predicted[index],
                            "ok": ok,
                        })
                if compared:
                    score += 3 * matched / (compared + unverified)
                    validation = {
                        "matched": matched, "compared": compared, "unverified": unverified,
                        "calculationErrors": len(result["errors"]),
                        "maxError": _stat_text(max_error),
                        "samples": samples,
                    }
            except Exception:  # noqa: BLE001 反算失败不影响候选
                validation = None
        candidates.append({
            "formulaId": item.get("id") or "",
            "handwriting": item.get("handwriting"),
            "target": target,
            "expr": expr,
            "formula": formula,
            "mapping": mapping,
            "units": {var: units.get(var, "") for var in mapping},
            "unitConflicts": unit_conflicts,
            "columnKey": target_key,
            "score": round(min(10.0, max(0.0, score)), 2),
            "validation": validation,
        })
    candidates.sort(key=lambda c: c["score"], reverse=True)
    return candidates[:3]


async def match_formulas(request: Request) -> JSONResponse:
    """从已保存公式库给当前表格推荐 Top-3 候选;不落库。"""
    rec = require_auth(request)
    uid = str(rec.get("uid") or "")
    table_id = str(request.path_params.get("table_id") or "")
    table = await asyncio.to_thread(_get_table, uid, table_id)
    if table is None:
        raise HTTPException(404, "not found")
    body = await read_json_object(request)
    doc_id = str(body.get("docId") or "").strip()
    saved = await asyncio.to_thread(
        rc._list_extracted_formulas, uid, table_id, doc_id)
    candidates = await asyncio.to_thread(_match_formula_candidates, table, saved)
    return JSONResponse({"candidates": candidates},
                        headers={"Cache-Control": "no-store"})


def _parse_topic(text: str) -> dict:
    raw = text.strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    start, end = raw.find("{"), raw.rfind("}")
    if start < 0 or end <= start:
        raise ValueError("模型未返回 JSON")
    data = json.loads(raw[start:end + 1])
    title = str(data.get("title") or "").strip().replace("\n", " ")[:40]
    return {"title": title, "confidence": rc._confidence(data.get("confidence"))}


async def suggest_topic(request: Request) -> JSONResponse:
    """根据表头、样例数据与来源论文建议表格主题;不落库。"""
    rec = require_auth(request)
    uid = str(rec.get("uid") or "")
    table_id = str(request.path_params.get("table_id") or "")
    table = await asyncio.to_thread(_get_table, uid, table_id)
    if table is None:
        raise HTTPException(404, "not found")
    if not llm_mod.is_configured():
        raise HTTPException(503, "模型未配置")
    if not _topic_limiter.allow(uid):
        raise HTTPException(429, "主题建议过于频繁,请稍后再试")
    columns = table.get("columns") or []
    labels = [str(c.get("label") or c.get("key") or "") for c in columns][:30]
    rows = table.get("rows") or []
    sample = [" | ".join(str(v) for v in row[:30]) for row in rows[:8]]
    source = table.get("source") or {}
    context = (
        f"表名: {table.get('title') or ''}\n"
        f"来源论文: {source.get('docName') or ''}\n"
        f"列: {'、'.join(labels)}\n"
        "样例数据:\n" + "\n".join(sample)
    )
    system = (
        "你是科研数据表命名助手。根据列名、样例数据和来源论文,判断这张表的简短主题。"
        "只输出 JSON,不要解释。格式:"
        '{"title":"不超过20字的主题","confidence":0到1之间的小数}。'
        "主题必须是具体研究对象或实验主题,不要用“数据表”“实验结果”这类空泛名称。"
        "信息不足时把 title 留空、confidence 设为 0。"
    )
    parts: list[str] = []
    async for delta in llm_mod.stream_chat(
            system, context, role="main", fallback=True, max_tokens=200):
        parts.append(delta)
    try:
        result = _parse_topic("".join(parts))
    except (ValueError, json.JSONDecodeError) as exc:
        raise HTTPException(502, f"主题建议失败: {exc}") from exc
    return JSONResponse(result, headers={"Cache-Control": "no-store"})


async def stats_table(request: Request) -> JSONResponse:
    """统计面板:不落库,直接分析当前草稿(含未保存修改)。"""
    rec = require_auth(request)
    uid = str(rec.get("uid") or "")
    table_id = str(request.path_params.get("table_id") or "")
    if await asyncio.to_thread(_get_table, uid, table_id) is None:
        raise HTTPException(404, "not found")
    body = await read_json_object(request)
    try:
        (columns, rows, formulas, _row_labels, _merges,
         _col_widths, _row_heights, _row_sources) = _validate_table_payload(body)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    raw_labels = body.get("rowLabels")
    row_labels = ([str(v)[:60] for v in raw_labels[:len(rows)]]
                  if isinstance(raw_labels, list) else [])
    result = await asyncio.to_thread(
        analyze, columns, rows, formulas, row_labels)
    return JSONResponse(result, headers={"Cache-Control": "no-store"})


async def export_table(request: Request) -> Response:
    rec = require_auth(request)
    uid = str(rec.get("uid") or "")
    table_id = str(request.path_params.get("table_id") or "")
    fmt = (request.query_params.get("format") or "csv").lower()
    table = await asyncio.to_thread(_get_table, uid, table_id)
    if table is None:
        raise HTTPException(404, "not found")
    if fmt == "xlsx":
        data = await asyncio.to_thread(_export_xlsx, table)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        data = await asyncio.to_thread(_export_csv, table)
        media = "text/csv; charset=utf-8"
    filename = f"{table['title']}.{fmt}"
    return Response(data, media_type=media, headers={
        "Content-Disposition": f"attachment; filename=\"table.{fmt}\"; "
                               f"filename*=UTF-8''{quote(filename)}",
        "Cache-Control": "no-store",
    })


async def print_table(request: Request) -> JSONResponse:
    """保存当前草稿并生成 xlsx,用本机 WPS/Excel 打开,由用户调整后打印。"""
    rec = require_auth(request)
    uid = str(rec.get("uid") or "")
    table_id = str(request.path_params.get("table_id") or "")
    if not _print_allowed(request):
        raise HTTPException(403, "打印仅支持本机访问")
    body = await read_json_object(request)
    try:
        saved = await asyncio.to_thread(_save_table, uid, table_id, body)
    except KeyError as exc:
        raise HTTPException(404, "not found") from exc
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    path = await asyncio.to_thread(_write_print_copy, uid, saved)
    opened = await asyncio.to_thread(_open_in_office, path)
    return JSONResponse({"ok": True, "opened": opened, "file": path.name})
