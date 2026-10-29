"""多文献对比表与项目报告导出。"""
from __future__ import annotations

import io

import pytest
from docx import Document

import research_compare as rcmp
import research_corpus as rc
import research_report as rrep
import research_tables as rt


@pytest.fixture()
def research_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(rc, "DATA_DIR", tmp_path / "research")
    rc._indexes.clear()
    rc._index_paths.clear()
    yield tmp_path


def test_compare_endpoint_builds_rows_and_sources(research_dir, monkeypatch):
    from starlette.applications import Starlette
    from starlette.routing import Route
    from starlette.testclient import TestClient

    monkeypatch.setattr(rcmp, "require_auth", lambda request: {"uid": "u-cmp"})
    monkeypatch.setattr(rcmp.llm, "is_configured", lambda: True)
    monkeypatch.setattr(rcmp, "_limiter", rcmp.SlidingWindowLimiter(100, 60))
    monkeypatch.setattr(rcmp, "_doc_meta", lambda uid, doc_id: {
        "id": doc_id, "name": f"{doc_id}.pdf", "title": f"论文 {doc_id}",
        "authors": "A. Author", "year": 2024, "venue": "Journal", "pages": 9,
    })
    monkeypatch.setattr(rcmp, "_excerpt", lambda uid, doc_id: ("正文", 3))

    async def fake_ask(meta, text):
        return {
            "question": f"{meta['title']} 的问题",
            "method": "方法", "data": "样本 30", "results": "提升 5%",
            "limitations": "",
        }

    monkeypatch.setattr(rcmp, "_ask_doc", fake_ask)
    app = Starlette(routes=[
        Route("/api/research/literature/compare", rcmp.compare, methods=["POST"]),
    ])
    client = TestClient(app)
    res = client.post("/api/research/literature/compare", json={"docIds": ["d1", "d2"]})
    assert res.status_code == 200, res.text
    body = res.json()
    assert len(body["rows"]) == 2
    assert body["columns"][0]["key"] == "A"
    assert body["rows"][0][0] == "论文 d1"
    assert body["rowSources"][1]["docId"] == "d2"
    assert body["rowSources"][1]["page"] == 3


def test_verify_numbers_flags_values_not_in_text():
    fields = {"data": "样本 30", "results": "提升 8.4 分,通过率 82%"}
    assert rcmp._verify_numbers(fields, "样本 30,提升 8.4 分,通过率 82%") == ""
    note = rcmp._verify_numbers(fields, "样本 30,提升 8.4 分")
    assert "82" in note and "未在原文找到" in note


def test_compare_requires_two_docs(research_dir, monkeypatch):
    from starlette.applications import Starlette
    from starlette.routing import Route
    from starlette.testclient import TestClient

    monkeypatch.setattr(rcmp, "require_auth", lambda request: {"uid": "u-cmp"})
    monkeypatch.setattr(rcmp.llm, "is_configured", lambda: True)
    app = Starlette(routes=[
        Route("/api/research/literature/compare", rcmp.compare, methods=["POST"]),
    ])
    client = TestClient(app)
    res = client.post("/api/research/literature/compare", json={"docIds": ["d1"]})
    assert res.status_code == 400


def test_report_docx_contains_project_tables_and_stats(research_dir):
    project = rc._create_project("u-rep", "研究X")
    table = rt._create_table("u-rep", {
        "title": "测量结果",
        "columns": [
            {"key": "A", "label": "x", "type": "number"},
            {"key": "B", "label": "2x", "type": "computed"},
        ],
        "rows": [["1", ""], ["2", ""], ["3", ""]],
        "formulas": {"B": "=A2*2"},
        "source": {"docId": "d1", "docName": "论文.pdf", "page": 4},
    }, project["id"])
    project["goal"] = "验证假设"
    data = rrep._build_docx("u-rep", project, [rt._get_table("u-rep", table["id"])])
    assert data.startswith(b"PK")
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "研究目标" in text and "验证假设" in text
    assert "数据与结果" in text and "测量结果" in text
    assert "描述性统计" in text
    assert "论文.pdf" in text
    assert doc.tables and len(doc.tables[0].rows) == 4


def test_report_endpoint_returns_docx(research_dir, monkeypatch):
    from starlette.applications import Starlette
    from starlette.routing import Route
    from starlette.testclient import TestClient

    project = rc._create_project("u-rep2", "报告X")
    monkeypatch.setattr(rrep, "require_auth", lambda request: {"uid": "u-rep2"})
    app = Starlette(routes=[
        Route("/api/research/report", rrep.export_report, methods=["POST"]),
    ])
    client = TestClient(app)
    res = client.post("/api/research/report", json={"projectId": project["id"]})
    assert res.status_code == 200, res.text
    assert res.content.startswith(b"PK")
    assert "wordprocessingml" in res.headers["content-type"]


def test_report_unknown_project_404(research_dir, monkeypatch):
    from starlette.applications import Starlette
    from starlette.routing import Route
    from starlette.testclient import TestClient

    monkeypatch.setattr(rrep, "require_auth", lambda request: {"uid": "u-rep3"})
    app = Starlette(routes=[
        Route("/api/research/report", rrep.export_report, methods=["POST"]),
    ])
    client = TestClient(app)
    res = client.post("/api/research/report", json={"projectId": "nope"})
    assert res.status_code == 404
