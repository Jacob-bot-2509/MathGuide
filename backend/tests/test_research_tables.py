"""实验数据表:公式引擎、模板导入、API、导出。"""
from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

import research_corpus as rc
import research_tables as rt


@pytest.fixture()
def tables(tmp_path, monkeypatch):
    monkeypatch.setattr(rc, "DATA_DIR", tmp_path / "research")
    rc._indexes.clear()
    rc._index_paths.clear()
    yield tmp_path


COLS = [
    {"key": "A", "label": "时间", "type": "number"},
    {"key": "B", "label": "位移", "type": "number"},
    {"key": "C", "label": "速度", "type": "computed"},
]


def test_column_letters_roundtrip():
    assert rt._col_to_index("A") == 0
    assert rt._col_to_index("Z") == 25
    assert rt._col_to_index("AA") == 26
    assert rt._index_to_col(0) == "A"
    assert rt._index_to_col(25) == "Z"
    assert rt._index_to_col(26) == "AA"


def test_formula_per_row_and_four_decimals():
    rows = [["1", "2"], ["3", "4"]]
    result = rt.compute(
        [{"key": "A", "label": "a", "type": "number"},
         {"key": "B", "label": "b", "type": "number"},
         {"key": "C", "label": "c", "type": "computed"}],
        rows, {"C": "=A2/B2"})
    assert result["computed"]["C"] == ["0.5000", "0.7500"]
    assert result["errors"] == []


def test_rounding_half_up_4dp():
    result = rt.compute(
        [{"key": "A", "label": "a", "type": "number"},
         {"key": "B", "label": "b", "type": "computed"}],
        [["2"]], {"B": "=1/A2"})
    assert result["computed"]["B"] == ["0.5000"]


def test_absolute_range_is_anchored():
    result = rt.compute(
        [{"key": "A", "label": "a", "type": "number"},
         {"key": "B", "label": "total", "type": "computed"}],
        [["1"], ["2"], ["3"]], {"B": "=SUM($A$2:$A$4)"})
    assert result["computed"]["B"] == ["6.0000"] * 3


def test_relative_range_shifts_per_row():
    result = rt.compute(
        [{"key": "A", "label": "a", "type": "number"},
         {"key": "B", "label": "sum", "type": "computed"}],
        [["1"], ["2"], ["3"]], {"B": "=SUM(A2:A3)"})
    assert result["computed"]["B"][:2] == ["3.0000", "5.0000"]


def test_division_by_zero_reports_cell_error():
    result = rt.compute(
        [{"key": "A", "label": "a", "type": "number"},
         {"key": "B", "label": "b", "type": "computed"}],
        [["1"], ["0"]], {"B": "=A2/0"})
    assert result["computed"]["B"][0] == ""
    assert result["errors"][0]["row"] == 0


def test_circular_dependency_rejected():
    result = rt.compute(
        [{"key": "A", "label": "a", "type": "computed"},
         {"key": "B", "label": "b", "type": "computed"}],
        [["1"]], {"A": "=B2", "B": "=A2"})
    assert result["computed"] == {}
    assert result["errors"]


def test_exponent_guard():
    result = rt.compute(
        [{"key": "A", "label": "a", "type": "number"},
         {"key": "B", "label": "b", "type": "computed"}],
        [["9"]], {"B": "=A2^999999"})
    assert result["computed"]["B"][0] == ""
    assert "指数" in result["errors"][0]["error"]


def test_no_data_cells_stay_blank_instead_of_crashing():
    result = rt.compute(
        [{"key": "A", "label": "a", "type": "number"},
         {"key": "B", "label": "b", "type": "number"},
         {"key": "C", "label": "c", "type": "computed"},
         {"key": "D", "label": "d", "type": "computed"}],
        [["", ""], ["2", ""]],
        {"C": "=A2/B2", "D": "=SUM(A2:A2)"})
    assert result["computed"]["C"] == ["", ""]
    assert result["computed"]["D"] == ["", "2.0000"]
    assert result["errors"]
    # 单格失败不影响同表其他行
    assert result["computed"]["D"][1] == "2.0000"


def test_finish_computes_every_formula_column():
    result = rt.compute(
        [{"key": "A", "label": "a", "type": "number"},
         {"key": "B", "label": "b", "type": "number"},
         {"key": "C", "label": "c", "type": "computed"},
         {"key": "D", "label": "d", "type": "computed"}],
        [["1", "2", ""], ["3", "4", ""]],
        {"C": "=A2/B2", "D": "=A2+B2"})
    assert result["computed"]["C"] == ["0.5000", "0.7500"]
    assert result["computed"]["D"] == ["3.0000", "7.0000"]
    assert result["errors"] == []


def test_delete_table_hard_removes_row(tables):
    table = rt._create_table("u-del", {
        "title": "待清除",
        "columns": COLS,
        "rows": [["1", "2", ""]],
        "formulas": {"C": "=A2/B2"},
    }, None)
    assert rt._delete_table("u-del", table["id"], hard=True) is True
    assert rt._get_table("u-del", table["id"]) is None
    assert rt._delete_table("u-del", table["id"], hard=True) is False


def test_merge_labels_widths_roundtrip(tables):
    table = rt._create_table("u-1", {
        "title": "m",
        "columns": [
            {"key": "A", "label": "时间", "type": "number"},
            {"key": "B", "label": "位移", "type": "number"},
        ],
        "rows": [["1", "2"], ["3", "4"]],
        "formulas": {},
        "rowLabels": ["第一组", "第二组"],
        "merges": [{"r1": 0, "c1": 0, "r2": 1, "c2": 0}],
        "colWidths": [150, 90],
        "rowHeights": [40, 55],
    }, None)
    loaded = rt._get_table("u-1", table["id"])
    assert loaded["rowLabels"] == ["第一组", "第二组"]
    assert loaded["merges"] == [{"r1": 0, "c1": 0, "r2": 1, "c2": 0}]
    assert loaded["colWidths"] == [150, 90]
    assert loaded["rowHeights"] == [40, 55]

    headers, rows = rt._table_rows_with_computed(loaded)
    assert headers[0] == "名称" and headers[1] == "时间"
    assert rows[0][0] == "第一组"

    data = rt._export_xlsx(loaded)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert "B2:B3" in {str(rng) for rng in ws.merged_cells.ranges}


def test_source_roundtrip_and_preserve_on_autosave(tables):
    source = {"docId": "doc-1", "docName": "paper.pdf", "page": 3, "quote": "Table 1"}
    table = rt._create_table("u-src", {
        "title": "转录",
        "columns": [{"key": "A", "label": "x", "type": "number"}],
        "rows": [["1"]],
        "formulas": {},
        "source": source,
    }, None)
    loaded = rt._get_table("u-src", table["id"])
    assert loaded["source"]["docId"] == "doc-1"
    assert loaded["source"]["page"] == 3

    # 旧客户端不带 source:自动保存不得抹掉来源
    rt._save_table("u-src", table["id"], {
        "title": "转录 2",
        "columns": [{"key": "A", "label": "x", "type": "number"}],
        "rows": [["2"]],
        "formulas": {},
    })
    loaded = rt._get_table("u-src", table["id"])
    assert loaded["title"] == "转录 2"
    assert loaded["source"]["docId"] == "doc-1"


def test_row_sources_roundtrip_and_preserve(tables):
    sources = [
        {"docId": "d1", "docName": "A.pdf", "page": 2},
        {"docId": "d2", "docName": "B.pdf", "page": 7},
    ]
    table = rt._create_table("u-rows", {
        "title": "对比",
        "columns": [{"key": "A", "label": "方法", "type": "text"}],
        "rows": [["方法一"], ["方法二"]],
        "formulas": {},
        "rowSources": sources,
    }, None)
    loaded = rt._get_table("u-rows", table["id"])
    assert [s["docId"] for s in loaded["rowSources"]] == ["d1", "d2"]
    assert loaded["rowSources"][1]["page"] == 7

    rt._save_table("u-rows", table["id"], {
        "title": "对比 2",
        "columns": [{"key": "A", "label": "方法", "type": "text"}],
        "rows": [["方法一"], ["方法二"]],
        "formulas": {},
    })
    loaded = rt._get_table("u-rows", table["id"])
    assert loaded["rowSources"][0]["docId"] == "d1"


def test_analyze_descriptive_stats_and_correlation():
    columns = [
        {"key": "A", "label": "x", "type": "number"},
        {"key": "B", "label": "y", "type": "number"},
        {"key": "C", "label": "double", "type": "computed"},
    ]
    rows = [["1", "2"], ["2", "4"], ["3", "6"], ["4", "8"]]
    result = rt.analyze(columns, rows, {"C": "=A2*2"}, ["一", "二", "三", "四"])
    stats = {s["key"]: s for s in result["stats"]}
    assert stats["A"]["n"] == 4
    assert stats["A"]["mean"] == "2.5000"
    assert stats["A"]["median"] == "2.5000"
    assert stats["A"]["stdev"] == "1.2910"
    assert stats["A"]["q1"] == "1.7500"
    assert stats["A"]["q3"] == "3.2500"
    assert stats["C"]["mean"] == "5.0000"
    r = next(c for c in result["correlations"] if {c["x"], c["y"]} == {"A", "B"})
    assert r["r"] == "1.0000"
    series = {s["key"]: s for s in result["series"]}
    assert series["C"]["values"] == [2, 4, 6, 8]
    assert result["rowLabels"] == ["一", "二", "三", "四"]


def test_analyze_skips_text_and_counts_missing():
    columns = [
        {"key": "A", "label": "group", "type": "text"},
        {"key": "B", "label": "x", "type": "number"},
    ]
    rows = [["ctrl", "1"], ["ctrl", ""], ["drug", "3"]]
    result = rt.analyze(columns, rows, {})
    assert [s["key"] for s in result["stats"]] == ["B"]
    assert result["stats"][0]["n"] == 2
    assert result["stats"][0]["missing"] == 1


def test_context_text_contains_data_and_truncation():
    columns = [{"key": "A", "label": "x", "type": "number"}]
    rows = [[str(i)] for i in range(60)]
    table = {
        "title": "长表", "columns": columns, "rows": rows, "formulas": {},
        "rowLabels": [], "source": {"docId": "d", "docName": "p.pdf", "page": 2},
    }
    text = rt.context_text(table, max_rows=40)
    assert "长表" in text and "60 行" in text
    assert "| 40 |" in text and "| 41 |" not in text
    assert "仅列出前 40 行" in text
    assert "描述性统计" in text and "均值=29.5000" in text
    assert "p.pdf" in text and "第 2 页" in text


def test_stats_api_rejects_unknown_table(tables, monkeypatch):
    from starlette.applications import Starlette
    from starlette.routing import Route
    from starlette.testclient import TestClient

    monkeypatch.setattr(rt, "require_auth", lambda request: {"uid": "u-stats"})
    app = Starlette(routes=[
        Route("/api/research/tables/{table_id}/stats", rt.stats_table, methods=["POST"]),
    ])
    client = TestClient(app)
    res = client.post("/api/research/tables/nope/stats", json={
        "columns": [{"key": "A"}], "rows": [["1"]], "formulas": {},
    })
    assert res.status_code == 404


def test_stats_api_returns_stats_for_saved_table(tables, monkeypatch):
    from starlette.applications import Starlette
    from starlette.routing import Route
    from starlette.testclient import TestClient

    monkeypatch.setattr(rt, "require_auth", lambda request: {"uid": "u-stats"})
    app = Starlette(routes=[
        Route("/api/research/tables/{table_id}/stats", rt.stats_table, methods=["POST"]),
    ])
    client = TestClient(app)
    created = rt._create_table("u-stats", {
        "title": "s",
        "columns": [
            {"key": "A", "label": "x", "type": "number"},
            {"key": "B", "label": "y", "type": "computed"},
        ],
        "rows": [["1", ""], ["2", ""], ["3", ""]],
        "formulas": {"B": "=A2*2"},
    }, None)
    res = client.post(f"/api/research/tables/{created['id']}/stats", json={
        "columns": [
            {"key": "A", "label": "x", "type": "number"},
            {"key": "B", "label": "y", "type": "computed"},
        ],
        "rows": [["1", ""], ["2", ""], ["3", ""]],
        "formulas": {"B": "=A2*2"},
        "rowLabels": ["a", "b", "c"],
    })
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["stats"][1]["mean"] == "4.0000"
    assert body["rowLabels"] == ["a", "b", "c"]
    assert body["computed"]["B"] == ["2.0000", "4.0000", "6.0000"]


def test_overlapping_merges_rejected(tables):
    with pytest.raises(ValueError):
        rt._create_table("u-1", {
            "columns": [{"key": "A"}, {"key": "B"}],
            "rows": [["", ""], ["", ""]],
            "formulas": {},
            "merges": [
                {"r1": 0, "c1": 0, "r2": 1, "c2": 0},
                {"r1": 1, "c1": 0, "r2": 1, "c2": 1},
            ],
        }, None)


def test_import_csv_and_xlsx(tables):
    csv_data = "t,x\n1,2\n3,4\n".encode("utf-8")
    table = rt._import_table("u-1", "exp.csv", csv_data, None)
    assert table["title"] == "exp"
    assert [c["key"] for c in table["columns"]] == ["A", "B"]
    assert table["rows"] == [["1", "2"], ["3", "4"]]

    wb = Workbook()
    ws = wb.active
    ws.append(["t", "x"])
    ws.append([1, 2])
    ws.append([3, 4])
    buf = io.BytesIO()
    wb.save(buf)
    table = rt._import_table("u-1", "exp.xlsx", buf.getvalue(), None)
    assert table["rows"] == [["1", "2"], ["3", "4"]]

    from docx import Document

    doc = Document()
    dt = doc.add_table(rows=2, cols=2)
    for r, values in enumerate((("t", "x"), ("5", "6"))):
        for c, value in enumerate(values):
            dt.cell(r, c).text = value
    dbuf = io.BytesIO()
    doc.save(dbuf)
    table = rt._import_table("u-1", "exp.docx", dbuf.getvalue(), None)
    assert table["rows"] == [["5", "6"]]


def test_table_api_flow(tables, monkeypatch):
    from starlette.applications import Starlette
    from starlette.routing import Route
    from starlette.testclient import TestClient

    monkeypatch.setattr(rt, "require_auth", lambda request: {"uid": "u-t"})
    app = Starlette(routes=[
        Route("/api/research/tables/import", rt.import_table, methods=["POST"]),
        Route("/api/research/tables", rt.list_tables),
        Route("/api/research/tables", rt.create_table, methods=["POST"]),
        Route("/api/research/tables/{table_id}", rt.save_table, methods=["POST"]),
        Route("/api/research/tables/{table_id}/compute", rt.compute_table, methods=["POST"]),
        Route("/api/research/tables/{table_id}/export", rt.export_table),
    ])
    client = TestClient(app)

    created = client.post("/api/research/tables", json={
        "title": "demo",
        "columns": COLS,
        "rows": [["1", "2", ""], ["3", "4", ""]],
        "formulas": {"C": "=A2/B2"},
    }).json()["table"]
    tid = created["id"]

    computed = client.post(f"/api/research/tables/{tid}/compute", json={
        "title": "demo",
        "columns": COLS,
        "rows": [["1", "2", ""], ["3", "4", ""]],
        "formulas": {"C": "=A2/B2"},
    }).json()
    assert computed["computed"]["C"] == ["0.5000", "0.7500"]

    csv_res = client.get(f"/api/research/tables/{tid}/export?format=csv")
    assert csv_res.status_code == 200
    assert "0.5000" in csv_res.text

    listed = client.get("/api/research/tables").json()["tables"]
    assert listed[0]["id"] == tid

    imported = client.post(
        "/api/research/tables/import?name=exp.csv",
        content="a,b\n1,2\n".encode("utf-8"),
        headers={"content-type": "text/csv"},
    )
    assert imported.status_code == 200
    assert imported.json()["table"]["rows"] == [["1", "2"]]


def test_print_writes_xlsx_and_opens(tables, monkeypatch):
    from starlette.applications import Starlette
    from starlette.routing import Route
    from starlette.testclient import TestClient

    monkeypatch.setattr(rt, "require_auth", lambda request: {"uid": "u-p"})
    opened: list = []
    monkeypatch.setattr(rt, "_open_in_office", lambda path: opened.append(path) or True)
    app = Starlette(routes=[
        Route("/api/research/tables/{table_id}/print", rt.print_table, methods=["POST"]),
    ])
    client = TestClient(app)
    payload = {
        "title": "打印测试",
        "columns": COLS,
        "rows": [["1", "2", ""], ["3", "4", ""]],
        "formulas": {"C": "=A2/B2"},
    }
    created = rt._create_table("u-p", payload, None)

    remote = client.post(
        f"/api/research/tables/{created['id']}/print", json=payload,
        headers={"host": "example.com"})
    assert remote.status_code == 403

    res = client.post(
        f"/api/research/tables/{created['id']}/print", json=payload,
        headers={"host": "127.0.0.1:8000"})
    assert res.status_code == 200
    body = res.json()
    assert body["opened"] is True and body["file"].endswith(".xlsx")
    assert opened and opened[0].exists()
    assert opened[0].read_bytes().startswith(b"PK")
