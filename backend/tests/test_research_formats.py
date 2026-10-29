"""多格式入库:按文件头识别并抽取正文,不依赖扩展名。"""
from __future__ import annotations

import io
import zipfile

import pytest

import research_corpus as rc


DOCX_TEXT = "Introduction\nWe study quantum error correction in trapped ions."
HTML_TEXT = "Introduction\nA short HTML paper about spectral gaps."
EPUB_TEXT = "Introduction\nAn EPUB chapter about modular forms."
RTF_TEXT = "Introduction\nAn RTF paper about Ricci flow."


@pytest.fixture()
def corpus(tmp_path, monkeypatch):
    monkeypatch.setattr(rc, "DATA_DIR", tmp_path / "research")
    monkeypatch.setattr(rc.embed_mod, "is_configured", lambda: False)
    rc._indexes.clear()
    rc._index_paths.clear()
    yield tmp_path
    rc._indexes.clear()
    rc._index_paths.clear()


def test_detect_format_by_content_not_extension():
    assert rc._detect_format("paper.txt", b"%PDF-1.7\n...", "text/plain") == "pdf"
    assert rc._detect_format("saved.pdf", b"<!doctype html><html><body>x</body></html>",
                             "application/pdf") == "html"
    assert rc._detect_format("note", b"{\\rtf1\\ansi hello}", "") == "rtf"
    assert rc._detect_format("paper.bin", "纯文本论文".encode("utf-8"), "") == "text"


def test_unknown_binary_rejected():
    with pytest.raises(ValueError):
        rc._detect_format("x.bin", b"\x00\x01\x02\x03binary", "application/octet-stream")


def test_docx_html_epub_rtf_ingest(corpus):
    from docx import Document

    doc = Document()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph(DOCX_TEXT.split("\n")[1])
    buf = io.BytesIO()
    doc.save(buf)
    docx_res = rc._ingest("u-fmt", "paper.docx", buf.getvalue(),
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert docx_res["chunks"] >= 1

    html_res = rc._ingest(
        "u-fmt", "paper.html",
        f"<html><body><h1>Introduction</h1><p>{HTML_TEXT.splitlines()[1]}</p></body></html>".encode(),
        "text/html")
    assert html_res["chunks"] >= 1

    epub = io.BytesIO()
    with zipfile.ZipFile(epub, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr("OEBPS/ch1.xhtml",
                    f"<html><body><h1>Introduction</h1><p>{EPUB_TEXT.splitlines()[1]}</p></body></html>")
    epub_res = rc._ingest("u-fmt", "paper.epub", epub.getvalue(), "application/epub+zip")
    assert epub_res["chunks"] >= 1

    rtf_res = rc._ingest(
        "u-fmt", "paper.rtf",
        ("{\\rtf1\\ansi\\deff0 {\\fonttbl {\\f0 Times;}}"
         "\\f0\\fs24 Introduction\\par An RTF paper about Ricci flow.}").encode(),
        "application/rtf")
    assert rtf_res["chunks"] >= 1

    index = rc._get_index("u-fmt")
    text = "\n".join(chunk.text for chunk in index.chunks)
    assert "quantum error correction" in text
    assert "spectral gaps" in text
    assert "modular forms" in text
    assert "Ricci flow" in text


def test_long_document_paginates_and_indexes(corpus):
    pages = rc._pack_pages("\n".join(
        f"Paragraph {i} about long proofs and spectral estimates." for i in range(12000)))
    assert len(pages) > 80
    text = "\n\n".join(t for _, t in pages)
    res = rc._ingest("u-long", "long.txt", text.encode("utf-8"), "text/plain")
    assert res["pages"] == len(pages)
    assert res["chunks"] >= 50
