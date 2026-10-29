"""项目研究报告导出:项目档案 + 数据表 + 服务端统计 + 来源,生成可编辑 DOCX。

不调用 LLM:报告内容全部来自用户已保存的项目与数据表,避免二次生成引入幻觉。
"""
from __future__ import annotations

import asyncio
import io
import re
import time
from urllib.parse import quote

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from starlette.exceptions import HTTPException
from starlette.requests import Request
from starlette.responses import JSONResponse, Response

import research_corpus as rc
import research_tables as rt
from auth import require_auth
from http_utils import read_json_object

STATUS_ZH = {
    "exploring": "探索中", "active": "进行中", "paused": "已暂停", "done": "已完成",
}
_CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _text(value: object, limit: int = 8000) -> str:
    return _CTRL_RE.sub("", str(value or ""))[:limit]


def _project(uid: str, project_id: str) -> dict | None:
    for project in rc._load_projects(uid):
        if project["id"] == project_id:
            return project
    return None


def _project_docs(uid: str, project_id: str) -> list[dict]:
    conn = rc._connect(uid)
    try:
        rows = conn.execute(
            "SELECT title, name, authors, year, venue, doi FROM docs"
            " WHERE project_id = ? AND deleted_at IS NULL ORDER BY created_at",
            (project_id,)).fetchall()
        return [
            {"title": r[0] or r[1], "authors": r[2] or "", "year": r[3],
             "venue": r[4] or "", "doi": r[5] or ""}
            for r in rows
        ]
    finally:
        conn.close()


def _unique_sources(tables: list[dict]) -> list[dict]:
    seen: set[str] = set()
    out: list[dict] = []
    for table in tables:
        candidates = []
        if isinstance(table.get("source"), dict):
            candidates.append(table["source"])
        for source in table.get("rowSources") or []:
            if isinstance(source, dict):
                candidates.append(source)
        for source in candidates:
            doc_id = str(source.get("docId") or "")
            if not doc_id or doc_id in seen:
                continue
            seen.add(doc_id)
            out.append(source)
    return out


def _stats_text(table: dict) -> list[str]:
    result = rt.analyze(
        table.get("columns") or [], table.get("rows") or [],
        table.get("formulas") or {}, table.get("rowLabels") or [])
    lines: list[str] = []
    for s in result["stats"]:
        parts = [f"n={s['n']}", f"均值={s['mean'] or '—'}",
                 f"中位数={s['median'] or '—'}", f"标准差={s['stdev'] or '—'}",
                 f"最小={s['min'] or '—'}", f"最大={s['max'] or '—'}"]
        lines.append(f"{s['key']} {s['label']}: " + ", ".join(parts))
    for c in result["correlations"]:
        lines.append(f"Pearson r({c['x']}, {c['y']}) = {c['r']}")
    return lines


def _add_paragraph(doc: Document, text: str, bold_prefix: str = "") -> None:
    if not text.strip():
        return
    p = doc.add_paragraph()
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
    p.add_run(_text(text))


def _add_table(doc: Document, table: dict) -> None:
    headers, rows = rt._table_rows_with_computed(table)
    if not headers:
        return
    grid = doc.add_table(rows=1, cols=len(headers))
    grid.style = "Table Grid"
    for i, label in enumerate(headers):
        cell = grid.rows[0].cells[i]
        cell.text = _text(label, 120)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9)
    for row in rows[:200]:
        cells = grid.add_row().cells
        for i, value in enumerate(row[:len(headers)]):
            cells[i].text = _text(value, 500)
            for paragraph in cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def _build_docx(uid: str, project: dict, tables: list[dict]) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "DengXian"
    style.font.size = Pt(10.5)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "等线")

    title = f"{project['title']} · 研究报告"
    doc.core_properties.title = title
    heading = doc.add_heading(title, level=0)
    heading.alignment = 0

    meta = doc.add_paragraph(
        f"状态:{STATUS_ZH.get(project['status'], project['status'])} · "
        f"生成日期:{time.strftime('%Y-%m-%d')} · MathGuide")
    meta.runs[0].font.size = Pt(9)

    doc.add_heading("研究目标", level=1)
    _add_paragraph(doc, project["goal"] or "（未填写）")
    doc.add_heading("研究思路", level=1)
    _add_paragraph(doc, project["approach"] or "（未填写）")
    doc.add_heading("过程笔记", level=1)
    _add_paragraph(doc, project["notes"] or "（未填写）")

    doc.add_heading("数据与结果", level=1)
    if not tables:
        doc.add_paragraph("（本项目还没有数据表）")
    for index, table in enumerate(tables, 1):
        doc.add_heading(f"{index}. {table.get('title') or '数据表'}", level=2)
        shape = f"{len(table.get('rows') or [])} 行 × {len(table.get('columns') or [])} 列"
        source = table.get("source") if isinstance(table.get("source"), dict) else {}
        source_text = ""
        if source.get("docName"):
            source_text = (f";来源:{_text(source['docName'], 200)}"
                           f"{' · 第 ' + str(source['page']) + ' 页' if source.get('page') else ''}")
        doc.add_paragraph(f"规模:{shape}{source_text}")

        formulas = table.get("formulas") or {}
        if formulas:
            doc.add_paragraph(
                "计算公式:" + ";".join(
                    f"{k} = {_text(v, 200).lstrip('=')}" for k, v in formulas.items()))
        _add_table(doc, table)

        stats = _stats_text(table)
        if stats:
            p = doc.add_paragraph()
            run = p.add_run("描述性统计（服务端计算）")
            run.bold = True
            for line in stats:
                doc.add_paragraph(_text(line, 500))

    sources = _unique_sources(tables)
    docs = _project_docs(uid, project["id"])
    if docs:
        doc.add_heading("项目文献", level=1)
        for item in docs:
            line = " · ".join(x for x in [
                _text(item["title"], 300), _text(item["authors"], 200),
                str(item["year"] or ""), _text(item["venue"], 120),
                _text(item["doi"], 120),
            ] if x)
            doc.add_paragraph(line)
    if sources:
        doc.add_heading("数据来源", level=1)
        for source in sources:
            line = f"{_text(source.get('docName') or source.get('docId'), 300)}"
            if source.get("page"):
                line += f" · 第 {source['page']} 页"
            doc.add_paragraph(line)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def export_report(request: Request) -> Response | JSONResponse:
    rec = require_auth(request)
    uid = str(rec.get("uid") or "")
    body = await read_json_object(request)
    project_id = str(body.get("projectId") or "").strip()
    if not project_id:
        raise HTTPException(400, "缺少 projectId")
    project = await asyncio.to_thread(_project, uid, project_id)
    if project is None:
        raise HTTPException(404, "研究项目不存在")
    raw_ids = body.get("tableIds")
    wanted = ({str(v) for v in raw_ids if str(v)} if isinstance(raw_ids, list) else set())
    summaries = await asyncio.to_thread(rt._list_tables, uid, project_id)
    tables = []
    for summary in summaries:
        if wanted and summary["id"] not in wanted:
            continue
        full = await asyncio.to_thread(rt._get_table, uid, summary["id"])
        if full:
            tables.append(full)
    data = await asyncio.to_thread(_build_docx, uid, project, tables)
    filename = f"{project['title'] or '研究报告'}.docx"
    return Response(data, media_type=(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        headers={
            "Content-Disposition": f"attachment; filename=\"report.docx\"; "
                                   f"filename*=UTF-8''{quote(filename)}",
            "Cache-Control": "no-store",
        })
